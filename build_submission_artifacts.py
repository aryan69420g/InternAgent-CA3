from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from docx import Document
from docx.shared import Inches as DInches, Pt as DPt, RGBColor as DRGB
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parent
DOCS = ROOT / "docs"
DOCS.mkdir(exist_ok=True)

NAVY = (18, 36, 58)
BLUE = (63, 115, 241)
TEAL = (21, 168, 151)
PALE = (242, 246, 252)
INK = (31, 45, 64)
MUTED = (92, 108, 128)
WHITE = (255, 255, 255)
ORANGE = (235, 158, 67)

def font(size, bold=False):
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ]
    for p in candidates:
        if Path(p).exists():
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()

def draw_architecture():
    W, H = 1600, 1000
    im = Image.new("RGB", (W, H), (248, 250, 253))
    d = ImageDraw.Draw(im)
    d.rounded_rectangle((35, 30, W-35, H-30), radius=34, fill=WHITE, outline=(222, 230, 241), width=3)
    d.text((85, 65), "INTERNAGENT", font=font(25, True), fill=TEAL)
    d.text((85, 110), "Multi-agent internship matching workflow", font=font(42, True), fill=NAVY)
    d.text((85, 170), "Student profile → plan → data retrieval → eligibility and scoring → critic validation → results", font=font(22), fill=MUTED)

    def box(coords, title, body, color):
        d.rounded_rectangle(coords, radius=24, fill=WHITE, outline=color, width=4)
        x1,y1,x2,y2 = coords
        d.rounded_rectangle((x1, y1, x2, y1+58), radius=20, fill=color)
        d.rectangle((x1, y1+38, x2, y1+58), fill=color)
        d.text((x1+22, y1+13), title, font=font(22, True), fill=WHITE)
        lines = body.split("\n")
        for i, line in enumerate(lines):
            d.text((x1+22, y1+78+i*32), line, font=font(18), fill=INK)

    def arrow(a, b, color=BLUE):
        d.line((a[0], a[1], b[0], b[1]), fill=color, width=6)
        import math
        ang = math.atan2(b[1]-a[1], b[0]-a[0])
        size = 18
        p1 = (b[0] - size*math.cos(ang-0.55), b[1] - size*math.sin(ang-0.55))
        p2 = (b[0] - size*math.cos(ang+0.55), b[1] - size*math.sin(ang+0.55))
        d.polygon([b, p1, p2], fill=color)

    box((90, 280, 480, 440), "1 · Planner Agent", "Interprets the profile\nCreates ordered tasks", BLUE)
    box((600, 280, 990, 440), "2 · Research Agent", "Loads synthetic CSV\nFilters by stipend preference", TEAL)
    box((1110, 280, 1510, 440), "3 · Matching Agent", "Checks year and CGPA\nScores fit; finds skill gaps", BLUE)
    arrow((480,360),(600,360))
    arrow((990,360),(1110,360))

    box((1110, 570, 1510, 745), "4 · Critic Agent", "Recomputes score\nChecks rules and ranking", ORANGE)
    box((600, 570, 990, 745), "Validation Gate", "Pass → publish results\nFail → one repair pass", TEAL)
    box((90, 570, 480, 745), "Streamlit Dashboard", "Ranked roles and score chart\nTrace, watchlist, CSV/JSON", BLUE)
    arrow((1310,440),(1310,570))
    arrow((1110,655),(990,655), ORANGE)
    arrow((600,655),(480,655), TEAL)
    # One bounded feedback loop from validation gate back toward critic.
    d.arc((840, 465, 1210, 585), start=190, end=350, fill=ORANGE, width=5)
    d.text((868, 486), "Repair / revalidate once", font=font(16, True), fill=(167, 102, 23))
    d.text((90, 825), "Tools: Python functions · Pandas CSV read/filter · deterministic scoring · independent validation", font=font(19), fill=MUTED)
    d.text((90, 865), "Data are fictional sample records. No external API or paid LLM is required.", font=font(19), fill=MUTED)
    im.save(DOCS / "architecture.png", optimize=True)

def rgb(c):
    return RGBColor(*c)

def ppt_text(slide, x, y, w, h, text, size=18, color=INK, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Aptos"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = rgb(color)
    if align is not None:
        p.alignment = align
    return shape

def ppt_card(slide, x, y, w, h, title, body, accent=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(WHITE)
    shape.line.color.rgb = rgb((220, 228, 239))
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.09), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(accent)
    bar.line.fill.background()
    ppt_text(slide, x+0.22, y+0.14, w-0.36, 0.42, title, 17, NAVY, True)
    ppt_text(slide, x+0.22, y+0.60, w-0.38, h-0.68, body, 12.5, MUTED)
    return shape

def ppt_base(prs, title, subtitle, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(PALE)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
    band.fill.solid(); band.fill.fore_color.rgb = rgb(TEAL); band.line.fill.background()
    ppt_text(slide, 0.62, 0.36, 12.0, 0.52, title, 26, NAVY, True)
    if subtitle:
        ppt_text(slide, 0.64, 0.90, 12.0, 0.42, subtitle, 11.5, MUTED)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(7.08), Inches(12.1), Inches(0.012))
    line.fill.solid(); line.fill.fore_color.rgb = rgb((222, 230, 241)); line.line.fill.background()
    ppt_text(slide, 0.64, 7.12, 10.8, 0.20, "INTERNAGENT  ·  CA-3 AGENTIC AI PoC  ·  SYNTHETIC DATA", 8.5, MUTED)
    ppt_text(slide, 12.0, 7.11, 0.55, 0.22, str(number).zfill(2), 9, TEAL, True, PP_ALIGN.RIGHT)
    return slide

def make_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(NAVY)
    accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), prs.slide_height)
    accent.fill.solid(); accent.fill.fore_color.rgb = rgb(TEAL); accent.line.fill.background()
    ppt_text(s, 0.9, 0.8, 6.0, 0.4, "CA-3 · AGENTIC AI PROOF OF CONCEPT", 12, (157,229,218), True)
    ppt_text(s, 0.9, 1.55, 7.6, 1.25, "InternAgent", 42, WHITE, True)
    ppt_text(s, 0.94, 2.95, 7.2, 1.0, "Internship Matching &\nEligibility Assistant", 25, (224,234,246), True)
    ppt_text(s, 0.94, 4.35, 6.8, 0.85, "A transparent multi-agent workflow for planning, retrieval, matching and validation.", 16, (194,211,231))
    ppt_text(s, 0.94, 6.62, 8.1, 0.32, "Aryan Gawande  ·  PRN: 24070126223  ·  Education & Career Development", 10, (157,229,218), True)
    shape = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.0), Inches(1.55), Inches(2.7), Inches(2.7))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(BLUE); shape.line.fill.background()
    ppt_text(s, 9.22, 2.28, 2.25, 0.9, "PLAN\nMATCH\nVERIFY", 17, WHITE, True, PP_ALIGN.CENTER)
    ppt_text(s, 8.35, 4.7, 4.0, 0.7, "Python  ·  Streamlit  ·  Pandas\nNo paid API required", 14, (194,211,231), False, PP_ALIGN.CENTER)

    s = ppt_base(prs, "The problem — and our solution", "A small, specific workflow that saves repeated manual checking.", 2)
    ppt_card(s, 0.7, 1.65, 5.85, 3.9, "Today’s friction", "• Compare listings across multiple sources\n\n• Interpret academic year and CGPA requirements\n\n• Estimate skill fit and identify missing skills\n\n• Trust recommendations without transparent checks", ORANGE)
    ppt_card(s, 6.82, 1.65, 5.78, 3.9, "InternAgent workflow", "• Collect a profile and preferences\n\n• Retrieve candidates from a synthetic CSV\n\n• Apply hard eligibility rules before ranking\n\n• Explain scores and ask a Critic Agent to verify them", TEAL)
    ppt_text(s, 0.85, 6.0, 11.8, 0.55, "Outcome: a ranked, explainable shortlist with eligibility evidence and skill gaps.", 18, NAVY, True)

    s = ppt_base(prs, "Architecture and data flow", "Agents exchange structured outputs; the Critic can trigger one repair/revalidation pass.", 3)
    s.shapes.add_picture(str(DOCS / "architecture.png"), Inches(0.65), Inches(1.45), width=Inches(12.0), height=Inches(5.2))

    s = ppt_base(prs, "Agent responsibilities", "Separate stages with explicit inputs, outputs and validation responsibilities.", 4)
    ppt_card(s, 0.7, 1.55, 5.8, 1.8, "1 · Planner Agent", "Creates an ordered plan: profile → retrieval → eligibility → score → validate.", BLUE)
    ppt_card(s, 6.8, 1.55, 5.8, 1.8, "2 · Research Agent", "Loads the fictional CSV and applies the minimum-stipend preference.", TEAL)
    ppt_card(s, 0.7, 3.65, 5.8, 1.8, "3 · Matching & Eligibility Agent", "Applies year/CGPA hard gates, computes match components and surfaces skill gaps.", BLUE)
    ppt_card(s, 6.8, 3.65, 5.8, 1.8, "4 · Critic / Validation Agent", "Independently recomputes scores and checks eligibility, bounds, order, IDs and top-k.", ORANGE)
    ppt_text(s, 0.85, 6.0, 11.4, 0.55, "Orchestrator: runs stages in order and permits one bounded repair/revalidation pass.", 16, NAVY, True)

    s = ppt_base(prs, "Explainable scoring — eligibility is separate", "Score is a fit heuristic, not a probability of getting hired.", 5)
    components = [("55%", "Required skills", BLUE), ("20%", "Domain interest", TEAL), ("10%", "Location / work mode", ORANGE), ("15%", "CGPA buffer", NAVY)]
    for i, (pct, label, color) in enumerate(components):
        y = 1.6 + i*1.08
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(y), Inches(2.0), Inches(0.68))
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(color); sh.line.fill.background()
        ppt_text(s, 0.98, y+0.12, 1.7, 0.38, pct, 19, WHITE, True, PP_ALIGN.CENTER)
        ppt_text(s, 3.1, y+0.10, 6.0, 0.42, label, 17, NAVY, True)
    ppt_card(s, 9.0, 1.72, 3.3, 3.25, "Hard eligibility gate", "Academic year must be within the role range and CGPA must meet the minimum.\n\nMissing skills reduce ranking fit but do not automatically disqualify.", ORANGE)
    ppt_text(s, 0.9, 6.15, 11.8, 0.45, "Scores are clamped to 0–100; ranked roles include skill gaps and score components.", 14, MUTED)

    s = ppt_base(prs, "PoC demonstration and validation", "Default demo profile: Year 2 · CGPA 7.5 · Python / ML / Pandas / Scikit-learn / SQL", 6)
    ppt_card(s, 0.75, 1.55, 3.8, 2.2, "12 source records", "All listings and organization names are synthetic.", BLUE)
    ppt_card(s, 4.78, 1.55, 3.8, 2.2, "Ranked shortlist", "Up to three eligible results with explanations and skill gaps.", TEAL)
    ppt_card(s, 8.82, 1.55, 3.8, 2.2, "Critic 5 / 5", "Five checks validate the returned recommendations.", ORANGE)
    ppt_text(s, 0.85, 4.25, 11.7, 0.5, "Six unit tests passed in the authoring environment.", 20, NAVY, True)
    ppt_text(s, 0.88, 4.95, 11.5, 1.0, "Live demo sequence: run workflow → inspect score chart → expand a role → show validation and execution trace → inspect ineligible roles → download CSV/JSON.", 15, MUTED)

    s = ppt_base(prs, "Limitations, future scope and close", "Built to demonstrate meaningful orchestration without unnecessary complexity.", 7)
    ppt_card(s, 0.72, 1.55, 5.85, 3.6, "Current limitations", "• Synthetic and small dataset\n\n• Exact normalized skill labels\n\n• Simplified academic rules\n\n• No live postings or application submission", ORANGE)
    ppt_card(s, 6.8, 1.55, 5.85, 3.6, "Future scope", "• Verified, refreshed internship sources\n\n• Semantic skill matching\n\n• LangGraph orchestration\n\n• Local-model explanations and application tracking", TEAL)
    ppt_text(s, 1.0, 5.75, 11.2, 0.65, "Agentic value = planning + tool use + decisions + independent verification.", 21, NAVY, True, PP_ALIGN.CENTER)
    ppt_text(s, 1.0, 6.45, 11.2, 0.35, "Thank you · Questions?", 16, TEAL, True, PP_ALIGN.CENTER)

    prs.save(ROOT / "CA3_Agentic_AI_Presentation.pptx")

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def docx_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, head in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = head
        set_cell_shading(cell, "12243A")
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = DRGB(255,255,255)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()
    return table

def make_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = DInches(0.72)
    sec.bottom_margin = DInches(0.68)
    sec.left_margin = DInches(0.8)
    sec.right_margin = DInches(0.8)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = DPt(10)
    styles["Normal"].font.color.rgb = DRGB(*INK)
    for name, size, color in [("Title", 30, NAVY), ("Heading 1", 19, NAVY), ("Heading 2", 13, BLUE), ("Heading 3", 11, TEAL)]:
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = DPt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = DRGB(*color)
    header = sec.header.paragraphs[0]
    header.text = "INTERNAGENT  |  CA-3 AGENTIC AI PROOF OF CONCEPT"
    header.style = styles["Caption"]
    header.runs[0].font.color.rgb = DRGB(*TEAL)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("InternAgent · Synthetic data only · ").font.color.rgb = DRGB(*MUTED)
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CA-3 PROJECT REPORT")
    r.bold = True; r.font.size = DPt(13); r.font.color.rgb = DRGB(*TEAL)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("InternAgent")
    r.bold = True; r.font.size = DPt(34); r.font.color.rgb = DRGB(*NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Agentic AI Internship Matching & Eligibility Assistant")
    r.font.size = DPt(17); r.font.color.rgb = DRGB(*BLUE)
    doc.add_paragraph()
    doc.add_picture(str(DOCS / "architecture.png"), width=DInches(6.35))
    p = doc.add_paragraph("Submitted by: Aryan Gawande\nPRN: 24070126223\nDomain: Education & Career Development\nTechnology: Python · Streamlit · Pandas · Plotly\nDate: 19 October 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.size = DPt(11); run.font.color.rgb = DRGB(*MUTED)
    doc.add_page_break()

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph("InternAgent is a working proof of concept for an explainable, multi-stage agentic workflow. It takes a student profile, retrieves candidate opportunities from a synthetic CSV dataset, applies academic eligibility rules, ranks eligible roles, exposes skill gaps, and validates recommendations before presentation. The app deliberately avoids paid APIs so the decision process remains transparent and reproducible.")
    doc.add_heading("2. Problem Statement and Domain Justification", level=1)
    doc.add_paragraph("Students often compare internship listings manually, interpret different academic requirements, and attempt to map personal skills to job requirements. This is repetitive, and a recommendation without an eligibility explanation or validation evidence may be difficult to trust. Education and career development is a suitable domain because the task naturally separates into planning, data retrieval, rule-based decision-making and output verification.")
    doc.add_heading("3. Existing Challenge", level=1)
    for item in [
        "Manually scanning several listings consumes time.",
        "Academic-year and CGPA constraints can be overlooked.",
        "An unranked list does not help students prioritize opportunities.",
        "Students need actionable information about missing required skills.",
        "A plain chatbot response does not prove that tool execution and validation occurred.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("4. Proposed Solution", level=1)
    doc.add_paragraph("The Streamlit UI accepts academic year, CGPA, skills, interests, preferred location, work mode, minimum stipend and desired result count. The orchestrator invokes four functional agents and returns ranked results only after validation.")
    docx_table(doc, ["Agent", "Responsibility", "Output"], [
        ("Planner Agent", "Decomposes the goal into ordered tasks.", "Execution plan"),
        ("Research Agent", "Loads the synthetic CSV and applies the stipend preference.", "Candidate records"),
        ("Matching & Eligibility Agent", "Applies year/CGPA gates, computes scores and identifies skill gaps.", "Assessed and ranked opportunities"),
        ("Critic / Validation Agent", "Recomputes scores and checks IDs, eligibility, score bounds, order and top-k count.", "Validation report"),
    ])
    doc.add_heading("5. Architecture and Data Flow", level=1)
    doc.add_picture(str(DOCS / "architecture.png"), width=DInches(6.35))
    doc.add_paragraph("The orchestrator first creates a plan, then filters the source records, assesses eligibility and fit, and asks the Critic to validate the result. If a validation check fails, it allows one repair/revalidation pass. The Streamlit UI displays the ranked shortlist, execution trace and excluded roles.")
    doc.add_heading("6. Match Scoring and Eligibility", level=1)
    doc.add_paragraph("Eligibility is a hard rule: the student's academic year must be within the internship's year range and CGPA must be at least the minimum. Skill matching is a ranking signal rather than a hard gate, making skill gaps visible without silently excluding every opportunity that requires growth.")
    docx_table(doc, ["Score component", "Weight", "Interpretation"], [
        ("Required skill overlap", "55%", "Share of required skills currently listed in the student profile"),
        ("Domain interest", "20%", "Higher score when the listing domain is one of the student's interests"),
        ("Location/work-mode fit", "10%", "Preference fit for location and work arrangement"),
        ("CGPA buffer", "15%", "Bounded reward for CGPA above the minimum requirement"),
    ])
    doc.add_paragraph("The total is clamped to the 0–100 range and sorted descending. The score is a transparent fit heuristic—not a probability of interview or hiring success.")
    doc.add_heading("7. Technology Stack", level=1)
    docx_table(doc, ["Technology", "Purpose"], [
        ("Python", "Workflow orchestration and rule functions"),
        ("Streamlit", "Interactive student profile and results dashboard"),
        ("Pandas / CSV", "Synthetic dataset loading and filtering"),
        ("Plotly", "Match-score visualization"),
        ("unittest", "Repeatable validation and regression tests"),
        ("GitHub Actions", "Runs tests on repository pushes and pull requests"),
    ])
    doc.add_heading("8. Output Validation and Testing", level=1)
    doc.add_paragraph("The Critic independently recomputes the weighted score without calling the matching helper. It checks uniqueness of internship IDs, eligibility of recommendations, score bounds and reproducibility, descending score order, and the top-k result-count policy.")
    docx_table(doc, ["Evidence", "Result"], [
        ("Synthetic dataset records", "12 fictional internship listings"),
        ("Automated unit tests", "6 tests passed in the authoring environment"),
        ("Default demo shortlist", "Up to 3 ranked recommendations"),
        ("Critic checks", "5 of 5 passed for the sample profile"),
    ])
    doc.add_heading("9. Results and Demonstration Plan", level=1)
    doc.add_paragraph("For the live demo, keep the sample profile (Year 2, CGPA 7.5, Python, Machine Learning, Pandas, Scikit-learn and SQL). Click Run agent workflow. Show the score chart, expand one role to review score components and missing skills, open the execution trace, inspect the eligibility watchlist and demonstrate CSV/JSON export.")
    doc.add_heading("10. Limitations", level=1)
    for item in [
        "The 12 listings and organization names are synthetic, not live internships.",
        "Skill match uses exact normalized labels and does not measure skill depth or semantic equivalence.",
        "Eligibility constraints are simplified and cannot represent every employer policy.",
        "The system does not submit applications or contact employers.",
        "The independent Critic improves reliability but cannot prove that every possible defect is absent.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("11. Future Scope", level=1)
    for item in [
        "Connect to a trusted, regularly refreshed source of internship postings.",
        "Add semantic skill matching and an optional local model for plain-language explanations.",
        "Represent agent state transitions explicitly with LangGraph.",
        "Add consent-based profile persistence, application tracking and feedback.",
        "Expand evaluation with adversarial profiles, rank-quality metrics and human review.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("12. Conclusion", level=1)
    doc.add_paragraph("InternAgent demonstrates a useful agentic pattern: a planner decomposes the task, a research stage uses a data tool, a matching stage makes rule-based decisions, and a Critic independently validates the result. The project emphasizes explainability and a real workflow over unnecessary complexity or a chatbot-only interface.")
    doc.add_heading("Reproducibility", level=1)
    doc.add_paragraph("Install the packages in requirements.txt, run streamlit run app.py, and run the automated tests with python -m unittest discover -s tests -v. All included data are synthetic and local.")
    doc.save(ROOT / "CA3_Technical_Report.docx")

if __name__ == "__main__":
    draw_architecture()
    make_ppt()
    make_docx()
    print("Generated architecture PNG, CA3_Agentic_AI_Presentation.pptx and CA3_Technical_Report.docx.")
